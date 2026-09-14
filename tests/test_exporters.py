import csv
import tempfile
import unittest
from pathlib import Path

from measurelog import exporters
from measurelog.db import Database
from measurelog.models import Location, Test


class ExportTestCase(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.out = Path(self.folder.name)
        self.db = Database(self.out / "test.db")
        self.location_a = self.db.save_location(Location(name="Tank A", code="TA", sort_order=0))
        self.location_b = self.db.save_location(Location(name="Tank B", code="TB", sort_order=1))
        self.ph = self.db.save_test(Test(name="pH", unit="", decimals=2, replicates=3,
                                         lower_limit=6.5, upper_limit=7.5, sort_order=0))
        self.cond = self.db.save_test(Test(name="Conductivity", unit="uS/cm", decimals=1,
                                           replicates=2, sort_order=1))
        self.run = self.db.create_run("2026-09-14", "08:00", "Round 1", "Sam")

        readings = [
            (self.location_a, self.ph, [7.01, 7.03, 6.99]),
            (self.location_a, self.cond, [412.0, 414.0]),
            (self.location_b, self.ph, [9.10, 9.20, 9.15]),   # above the upper limit
            (self.location_b, self.cond, [380.0, 384.0]),
        ]
        for location, test, values in readings:
            for replicate, value in enumerate(values, start=1):
                self.db.set_value(self.run.id, location.id, test.id, replicate, value)
        self.rows = self.db.fetch_long_rows()

    def tearDown(self):
        self.db.close()
        self.folder.cleanup()

    @staticmethod
    def read_csv(path):
        with open(path, newline="", encoding="utf-8-sig") as handle:
            return list(csv.reader(handle))


class RawExportTests(ExportTestCase):
    def test_one_row_per_replicate(self):
        path = exporters.export_raw_csv(self.out / "raw.csv", self.rows)
        table = self.read_csv(path)
        self.assertEqual(table[0], exporters.RAW_HEADERS)
        self.assertEqual(len(table) - 1, 10)

    def test_status_column_marks_readings_outside_the_limits(self):
        table = self.read_csv(exporters.export_raw_csv(self.out / "raw.csv", self.rows))
        status_index = exporters.RAW_HEADERS.index("Status")
        statuses = [row[status_index] for row in table[1:]]
        self.assertEqual(statuses.count("Above limit"), 3)
        self.assertEqual(statuses.count("In spec"), 3)
        self.assertEqual(statuses.count(""), 4)  # conductivity has no limits

    def test_notes_are_carried_through(self):
        self.db.set_value(self.run.id, self.location_a.id, self.ph.id, 1, 7.01, "re-read")
        table = self.read_csv(
            exporters.export_raw_csv(self.out / "raw.csv", self.db.fetch_long_rows()))
        self.assertIn("re-read", [row[exporters.RAW_HEADERS.index("Note")] for row in table[1:]])


class SummaryExportTests(ExportTestCase):
    def test_one_row_per_location_and_test(self):
        table = self.read_csv(exporters.export_summary_csv(self.out / "s.csv", self.rows))
        self.assertEqual(table[0], exporters.SUMMARY_HEADERS)
        self.assertEqual(len(table) - 1, 4)

    def test_statistics_are_correct(self):
        table = self.read_csv(exporters.export_summary_csv(self.out / "s.csv", self.rows))
        header = exporters.SUMMARY_HEADERS
        row = next(r for r in table[1:] if r[header.index("Location")] == "Tank A"
                   and r[header.index("Test")] == "pH")
        self.assertEqual(row[header.index("N")], "3")
        self.assertAlmostEqual(float(row[header.index("Mean")]), 7.01, places=4)
        self.assertAlmostEqual(float(row[header.index("SD")]), 0.02, places=4)
        self.assertAlmostEqual(float(row[header.index("%RSD")]), 0.29, places=2)
        self.assertAlmostEqual(float(row[header.index("Range")]), 0.04, places=4)
        self.assertEqual(row[header.index("Status")], "In spec")

    def test_out_of_spec_mean_is_flagged(self):
        table = self.read_csv(exporters.export_summary_csv(self.out / "s.csv", self.rows))
        header = exporters.SUMMARY_HEADERS
        row = next(r for r in table[1:] if r[header.index("Location")] == "Tank B"
                   and r[header.index("Test")] == "pH")
        self.assertEqual(row[header.index("Status")], "Above limit")

    def test_blank_statistics_export_as_empty_not_none(self):
        run = self.db.create_run("2026-09-15")
        self.db.set_value(run.id, self.location_a.id, self.ph.id, 1, 7.0)  # single value: no SD
        table = self.read_csv(
            exporters.export_summary_csv(self.out / "s.csv",
                                         self.db.fetch_long_rows(run_ids=[run.id])))
        self.assertEqual(table[1][exporters.SUMMARY_HEADERS.index("SD")], "")


class WideExportTests(ExportTestCase):
    def test_one_row_per_location_with_a_column_per_test(self):
        headers, table = exporters.wide_table(self.rows)
        self.assertEqual(headers[:5], ["Date", "Time", "Run", "Operator", "Location"])
        self.assertEqual(headers[5:], ["pH", "Conductivity (uS/cm)"])
        self.assertEqual(len(table), 2)

    def test_cells_hold_the_mean_of_the_replicates(self):
        _, table = exporters.wide_table(self.rows)
        row = next(r for r in table if r[4] == "Tank A")
        self.assertAlmostEqual(row[5], 7.01, places=4)
        self.assertAlmostEqual(row[6], 413.0, places=4)

    def test_missing_test_leaves_an_empty_cell(self):
        run = self.db.create_run("2026-09-15")
        self.db.set_value(run.id, self.location_a.id, self.ph.id, 1, 7.0)
        headers, table = exporters.wide_table(self.db.fetch_long_rows())
        row = next(r for r in table if r[0] == "2026-09-15")
        self.assertIsNone(row[headers.index("Conductivity (uS/cm)")])

    def test_written_file_has_blank_cells_not_the_word_none(self):
        run = self.db.create_run("2026-09-15")
        self.db.set_value(run.id, self.location_a.id, self.ph.id, 1, 7.0)
        path = exporters.export_wide_csv(self.out / "w.csv", self.db.fetch_long_rows())
        text = path.read_text(encoding="utf-8-sig")
        self.assertNotIn("None", text)


class ExcelExportTests(ExportTestCase):
    @unittest.skipUnless(exporters.excel_available(), "openpyxl is not installed")
    def test_workbook_has_all_four_sheets(self):
        import openpyxl

        path = exporters.export_excel(self.out / "book.xlsx", self.rows,
                                      self.db.list_locations(), self.db.list_tests())
        book = openpyxl.load_workbook(path)
        self.assertEqual(book.sheetnames, ["Raw data", "Summary", "Wide means", "Setup"])
        self.assertEqual(book["Raw data"].max_row, 11)
        self.assertEqual(book["Summary"].max_row, 5)
        self.assertEqual(book["Setup"].max_row, 5)

    @unittest.skipUnless(exporters.excel_available(), "openpyxl is not installed")
    def test_values_land_as_numbers_not_text(self):
        import openpyxl

        path = exporters.export_excel(self.out / "book.xlsx", self.rows)
        sheet = openpyxl.load_workbook(path)["Raw data"]
        value_column = exporters.RAW_HEADERS.index("Value") + 1
        self.assertIsInstance(sheet.cell(row=2, column=value_column).value, float)

    @unittest.skipUnless(exporters.excel_available(), "openpyxl is not installed")
    def test_empty_export_still_writes_a_valid_workbook(self):
        import openpyxl

        path = exporters.export_excel(self.out / "empty.xlsx", [])
        book = openpyxl.load_workbook(path)
        self.assertEqual(book["Raw data"].max_row, 1)


class HelperTests(ExportTestCase):
    def test_describe_counts_what_is_included(self):
        text = exporters.describe(self.rows)
        self.assertIn("10 measurements", text)
        self.assertIn("2 location(s)", text)
        self.assertIn("2 test(s)", text)

    def test_describe_handles_no_data(self):
        self.assertIn("No measurements", exporters.describe([]))

    def test_default_filename_includes_the_range(self):
        name = exporters.default_filename("measurelog", "csv", "2026-09-01", "2026-09-14")
        self.assertTrue(name.startswith("measurelog_2026-09-01_to_2026-09-14_"))
        self.assertTrue(name.endswith(".csv"))

    def test_default_filename_collapses_a_single_day(self):
        name = exporters.default_filename("m", "csv", "2026-09-14", "2026-09-14")
        self.assertIn("_2026-09-14_", name)
        self.assertNotIn("to", name)

    def test_write_csv_creates_missing_folders(self):
        path = exporters.write_csv(self.out / "deep" / "nested" / "f.csv", ["A"], [[1]])
        self.assertTrue(path.exists())

    def test_write_csv_reports_a_bad_path(self):
        blocker = self.out / "blocker"
        blocker.write_text("not a folder")
        with self.assertRaises(exporters.ExportError):
            exporters.write_csv(blocker / "impossible.csv", ["A"], [[1]])


if __name__ == "__main__":
    unittest.main()
