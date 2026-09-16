"""Read measurements out of a CSV or Excel file and into a run.

Two shapes of file turn up in practice, and both are accepted:

**Long** - one row per reading, with columns naming the location and the test::

    Location,Test,Replicate,Value
    Influent,pH,1,7.02

**Wide** - a grid, with one axis naming the tests and the other the
locations - the layout of the Review tab, of the "one row per location"
export, and of most hand-kept spreadsheets::

    Test,Influent,Digester 1,Effluent
    pH,7.02,7.11,7.26

Nothing is written until the import dialog has shown what was understood: this
module turns a file into a list of rows each carrying its own status, so a
misspelt test name is reported against the row it came from rather than
quietly dropping a reading.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

LAYOUT_LONG = "long"
LAYOUT_WIDE = "wide"

# Status values an ImportRow can carry. Only OK rows are written.
OK = "ok"
UNKNOWN_TEST = "unknown test"
UNKNOWN_LOCATION = "unknown location"
NOT_A_NUMBER = "not a number"
TOO_MANY_REPLICATES = "too many replicates"

MAX_PREVIEW_ROWS = 400

# Words that give a column away. An exact match beats a word appearing inside
# a longer header, so "Value" wins over "Measurement value" for the value role.
ROLE_HINTS: dict[str, tuple[str, ...]] = {
    "location": ("location", "site", "samplingpoint", "samplepoint", "point", "station",
                 "locationcode", "sample", "place", "tank"),
    "test": ("test", "parameter", "analyte", "determinand", "testcode", "analysis",
             "measurement", "property"),
    "value": ("value", "result", "reading", "measured", "amount", "concentration", "conc"),
    "replicate": ("replicate", "rep", "repeat", "duplicate"),
    "note": ("note", "notes", "comment", "comments", "remark", "remarks"),
}

# In a grid, these headings describe the round or restate a statistic; they are
# never a test or a location, so the importer steps over them.
METADATA_HEADINGS = frozenset({
    "date", "rundate", "time", "runtime", "run", "runlabel", "label", "operator",
    "note", "notes", "runnotes", "comment", "comments", "unit", "units", "id",
    "n", "mean", "sd", "rsd", "min", "max", "range", "status", "replicate",
    "lowerlimit", "upperlimit", "limits",
})

_TRAILING_BRACKET = re.compile(r"\s*\([^)]*\)\s*$")
_NON_ALNUM = re.compile(r"[^a-z0-9]+")
_LESS_OR_MORE = re.compile(r"^[<>≤≥]")
# Groups of exactly three digits mean the comma separates thousands, so "2,450"
# is 2450. "2,45" falls through and is read as the decimal 2.45.
_THOUSANDS = re.compile(r"^[+-]?\d{1,3}(,\d{3})+$")


class ImportFailed(RuntimeError):
    """The file could not be read at all."""


# --------------------------------------------------------------- name matching

def normalise(text: object) -> str:
    """Reduce a name to letters and digits, so spelling differences stop mattering."""
    return _NON_ALNUM.sub("", str(text or "").strip().lower())


def _keys(text: object) -> list[str]:
    """Every spelling of a name worth indexing: as written, and without units."""
    raw = str(text or "").strip()
    keys = [normalise(raw)]
    bare = _TRAILING_BRACKET.sub("", raw)
    if bare != raw:
        keys.append(normalise(bare))
    return [key for key in dict.fromkeys(keys) if key]


class Matcher:
    """Finds the test or location a spreadsheet heading is talking about.

    Matching is deliberately forgiving - "Total Solids (%)", "total solids" and
    "TotalSolids" all reach the same test - because a file written by hand or by
    an instrument rarely spells things the way the Setup tab does.
    """

    def __init__(self, items):
        self._exact: dict[str, object] = {}
        self._items = list(items)
        for item in self._items:
            for source in (getattr(item, "name", ""), getattr(item, "code", "")):
                for key in _keys(source):
                    self._exact.setdefault(key, item)

    def find(self, text: object):
        key = normalise(text)
        if not key:
            return None
        for candidate in _keys(text):
            if candidate in self._exact:
                return self._exact[candidate]
        # Fall back to a unique prefix match, which catches "Total Solids"
        # against "Total Solids dried". Ambiguity is left unmatched on purpose,
        # and short keys are excluded so the code "TS" cannot swallow "TSS".
        if len(key) < 3:
            return None
        hits = [item for stored, item in self._exact.items()
                if len(stored) >= 3 and (stored.startswith(key) or key.startswith(stored))]
        unique = {id(item): item for item in hits}
        return next(iter(unique.values())) if len(unique) == 1 else None


# -------------------------------------------------------------- reading files

@dataclass
class Table:
    """A sheet of cells, with its first useful row taken as the headings."""

    headers: list[str]
    rows: list[list]
    path: Path
    sheet: str = ""

    @property
    def width(self) -> int:
        return len(self.headers)

    def column(self, index: int) -> list:
        return [row[index] if index < len(row) else "" for row in self.rows]


def read_table(path: str | Path, sheet: str | None = None) -> Table:
    """Load a CSV, TSV or Excel file into headers and rows."""
    path = Path(path)
    if not path.is_file():
        raise ImportFailed(f"{path.name} could not be found.")
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return _read_excel(path, sheet)
    if suffix in (".csv", ".tsv", ".txt"):
        return _read_delimited(path)
    raise ImportFailed(
        f"MeasureLog cannot read {suffix or 'files without an extension'}. "
        "Save the data as CSV or as an Excel .xlsx workbook and try again."
    )


def sheet_names(path: str | Path) -> list[str]:
    """Worksheet names in a workbook; empty for anything else."""
    path = Path(path)
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        return []
    workbook = _open_workbook(path)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _open_workbook(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - only on a stripped build
        raise ImportFailed(
            "This build cannot read Excel files. Save the sheet as CSV and import that."
        ) from error
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise ImportFailed(f"{path.name} could not be opened: {error}") from error


def _read_excel(path: Path, sheet: str | None) -> Table:
    workbook = _open_workbook(path)
    try:
        worksheet = workbook[sheet] if sheet and sheet in workbook.sheetnames else workbook.active
        grid = [["" if cell is None else cell for cell in row]
                for row in worksheet.iter_rows(values_only=True)]
        name = worksheet.title
    finally:
        workbook.close()
    headers, rows = _split_header(grid)
    return Table(headers=headers, rows=rows, path=path, sheet=name)


def _read_delimited(path: Path) -> Table:
    text = _read_text(path)
    sample = text[:4096]
    if path.suffix.lower() == ".tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    grid = [row for row in csv.reader(text.splitlines(), delimiter=delimiter)]
    headers, rows = _split_header(grid)
    return Table(headers=headers, rows=rows, path=path)


def _read_text(path: Path) -> str:
    """Decode a text file, trying the encodings Excel actually writes."""
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
        except OSError as error:
            raise ImportFailed(f"{path.name} could not be read: {error}") from error
    raise ImportFailed(f"{path.name} is not a text file MeasureLog can read.")


def _split_header(grid: list[list]) -> tuple[list[str], list[list]]:
    """Take the first row with content as the headings; drop blank rows below."""
    index = next((i for i, row in enumerate(grid) if any(str(c).strip() for c in row)), None)
    if index is None:
        raise ImportFailed("That file is empty.")
    headers = [str(cell).strip() for cell in grid[index]]
    while headers and not headers[-1]:
        headers.pop()
    rows = [row for row in grid[index + 1:] if any(str(c).strip() for c in row)]
    return headers, rows


# ------------------------------------------------------------------- mapping

@dataclass
class Mapping:
    """Which column means what. The import dialog lets this be corrected."""

    layout: str = LAYOUT_LONG
    location: int | None = None
    test: int | None = None
    value: int | None = None
    replicate: int | None = None
    note: int | None = None
    # Wide layout only: the first column names one axis, the headings the other.
    key_column: int = 0
    key_role: str = "test"      # "test" when rows are tests, "location" when rows are locations
    skip: tuple[int, ...] = ()  # wide layout: columns that hold neither a test nor a location

    @property
    def complete(self) -> bool:
        """True when there is enough here to read the file without guessing."""
        if self.layout == LAYOUT_WIDE:
            return self.key_role in ("test", "location")
        needed = (self.location, self.test, self.value)
        # Three distinct columns: one column cannot name the location, name the
        # test and hold the reading all at once.
        return None not in needed and len(set(needed)) == 3


def guess_mapping(table: Table, tests, locations) -> Mapping:
    """Work out what the file looks like, so the dialog opens on the right answer."""
    wide = _guess_wide(table, tests, locations)
    long = _guess_long(table)
    # A file with a proper Location/Test/Value trio is unambiguous; otherwise a
    # row of headings that name known locations or tests means it is a grid.
    if long.complete:
        return long
    return wide if wide is not None else long


def _guess_long(table: Table) -> Mapping:
    scores: dict[str, list[tuple[int, int]]] = {role: [] for role in ROLE_HINTS}
    for index, header in enumerate(table.headers):
        key = normalise(header)
        if not key:
            continue
        for role, hints in ROLE_HINTS.items():
            if key in hints:
                scores[role].append((3, index))
            elif any(hint in key for hint in hints):
                scores[role].append((1, index))

    mapping = Mapping(layout=LAYOUT_LONG)
    taken: set[int] = set()
    # Strongest match first, so a clear "Value" column is not stolen by a role
    # that only matched it loosely.
    ranked = sorted(
        ((score, role, index) for role, hits in scores.items() for score, index in hits),
        key=lambda item: (-item[0], _ROLE_ORDER.index(item[1])),
    )
    for _score, role, index in ranked:
        if index in taken or getattr(mapping, role) is not None:
            continue
        setattr(mapping, role, index)
        taken.add(index)
    return mapping


_ROLE_ORDER = ["value", "location", "test", "replicate", "note"]


def _guess_wide(table: Table, tests, locations) -> Mapping | None:
    """Detect a grid by asking which column names one axis and the headings the other.

    The key column is not assumed to be the first: the app's own
    "one row per location" export carries Date, Time, Run and Operator in front
    of it, and those columns are marked to be stepped over.
    """
    test_matcher, location_matcher = Matcher(tests), Matcher(locations)
    best: tuple[int, Mapping] | None = None

    for key_index in range(min(table.width, 8)):
        headings = [(index, head) for index, head in enumerate(table.headers)
                    if index != key_index and str(head).strip()
                    and normalise(head) not in METADATA_HEADINGS]
        if len(headings) < 2:
            continue
        column = [cell for cell in table.column(key_index) if str(cell).strip()]
        if not column:
            continue

        head_locations = sum(1 for _, head in headings if location_matcher.find(head))
        head_tests = sum(1 for _, head in headings if test_matcher.find(head))
        column_tests = sum(1 for cell in column if test_matcher.find(cell))
        column_locations = sum(1 for cell in column if location_matcher.find(cell))

        for key_role, head_hits, column_hits in (
            ("test", head_locations, column_tests),
            ("location", head_tests, column_locations),
        ):
            if head_hits < 2:
                continue
            score = head_hits + column_hits
            if best is None or score > best[0]:
                skip = tuple(index for index in range(table.width)
                             if index != key_index
                             and index not in {i for i, _ in headings})
                best = (score, Mapping(layout=LAYOUT_WIDE, key_column=key_index,
                                       key_role=key_role, skip=skip))

    return best[1] if best else None


# -------------------------------------------------------------------- planning

@dataclass
class ImportRow:
    """One reading pulled out of the file, with what became of it."""

    source_row: int
    location_name: str = ""
    test_name: str = ""
    location_id: int | None = None
    test_id: int | None = None
    replicate: int = 1
    value: float | None = None
    note: str = ""
    status: str = OK
    detail: str = ""

    @property
    def ok(self) -> bool:
        return self.status == OK


@dataclass
class ImportPlan:
    """Everything the import dialog needs to show before anything is written."""

    rows: list[ImportRow] = field(default_factory=list)
    mapping: Mapping = field(default_factory=Mapping)
    table: Table | None = None

    @property
    def ready(self) -> list[ImportRow]:
        return [row for row in self.rows if row.ok]

    @property
    def problems(self) -> list[ImportRow]:
        return [row for row in self.rows if not row.ok]

    def summary(self) -> str:
        ready, problems = len(self.ready), len(self.problems)
        if not self.rows:
            return "No readings found in that file."
        text = f"{ready} reading{'s' if ready != 1 else ''} ready to import"
        if problems:
            counts: dict[str, int] = {}
            for row in self.problems:
                counts[row.status] = counts.get(row.status, 0) + 1
            listed = ", ".join(f"{count} {status}" for status, count in sorted(counts.items()))
            text += f"  •  {problems} skipped ({listed})"
        return text


def build_plan(table: Table, mapping: Mapping, tests, locations) -> ImportPlan:
    """Turn a table plus a mapping into rows, each one resolved or explained."""
    test_matcher, location_matcher = Matcher(tests), Matcher(locations)
    by_id = {test.id: test for test in tests}
    seen: dict[tuple[int | None, int | None], int] = {}

    if mapping.layout == LAYOUT_WIDE:
        rows = _wide_rows(table, mapping, test_matcher, location_matcher, seen)
    else:
        rows = _long_rows(table, mapping, test_matcher, location_matcher, seen)

    for row in rows:
        if row.ok:
            _check_replicate(row, by_id.get(row.test_id))
    return ImportPlan(rows=rows, mapping=mapping, table=table)


def _long_rows(table, mapping, test_matcher, location_matcher, seen) -> list[ImportRow]:
    rows: list[ImportRow] = []
    for offset, source in enumerate(table.rows):
        cell = lambda index: (source[index] if index is not None and index < len(source) else "")
        raw_value = cell(mapping.value)
        location_name = str(cell(mapping.location)).strip()
        test_name = str(cell(mapping.test)).strip()
        if not str(raw_value).strip() and not location_name and not test_name:
            continue
        row = ImportRow(
            source_row=offset + 2,  # +1 for the heading row, +1 for 1-based counting
            location_name=location_name,
            test_name=test_name,
            note=str(cell(mapping.note)).strip(),
        )
        _resolve(row, test_matcher, location_matcher)
        _set_value(row, raw_value)
        _set_replicate(row, cell(mapping.replicate), seen)
        rows.append(row)
    return rows


def _wide_rows(table, mapping, test_matcher, location_matcher, seen) -> list[ImportRow]:
    rows: list[ImportRow] = []
    key_index = mapping.key_column
    for offset, source in enumerate(table.rows):
        key_name = str(source[key_index]).strip() if key_index < len(source) else ""
        if not key_name:
            continue
        for index, heading in enumerate(table.headers):
            if index == key_index or index in mapping.skip or not str(heading).strip():
                continue
            raw_value = source[index] if index < len(source) else ""
            if not str(raw_value).strip():
                continue
            if mapping.key_role == "test":
                test_name, location_name = key_name, str(heading).strip()
            else:
                test_name, location_name = str(heading).strip(), key_name
            row = ImportRow(source_row=offset + 2, location_name=location_name,
                            test_name=test_name)
            _resolve(row, test_matcher, location_matcher)
            _set_value(row, raw_value)
            _set_replicate(row, "", seen)
            rows.append(row)
    return rows


def _resolve(row: ImportRow, test_matcher: Matcher, location_matcher: Matcher) -> None:
    test = test_matcher.find(row.test_name)
    location = location_matcher.find(row.location_name)
    row.test_id = getattr(test, "id", None)
    row.location_id = getattr(location, "id", None)
    if row.test_id is None:
        row.status, row.detail = UNKNOWN_TEST, (
            f"No test called \"{row.test_name}\"." if row.test_name
            else "No test named on this row."
        )
    elif row.location_id is None:
        row.status, row.detail = UNKNOWN_LOCATION, (
            f"No location called \"{row.location_name}\"." if row.location_name
            else "No location named on this row."
        )


def _set_value(row: ImportRow, raw) -> None:
    value, problem = to_number(raw)
    row.value = value
    if problem and row.status == OK:
        row.status, row.detail = NOT_A_NUMBER, problem


def _set_replicate(row: ImportRow, raw, seen: dict) -> None:
    """Use the file's replicate number, or count repeats of this test and location."""
    key = (row.location_id, row.test_id, row.location_name.lower(), row.test_name.lower())
    stated = str(raw).strip()
    if stated:
        try:
            row.replicate = max(1, int(float(stated.replace(",", "."))))
            seen[key] = max(seen.get(key, 0), row.replicate)
            return
        except ValueError:
            pass  # A label like "Rep A": fall through to counting.
    seen[key] = seen.get(key, 0) + 1
    row.replicate = seen[key]


def _check_replicate(row: ImportRow, test) -> None:
    allowed = max(1, int(getattr(test, "replicates", 1) or 1))
    if row.replicate > allowed:
        row.status = TOO_MANY_REPLICATES
        row.detail = (
            f"{row.test_name} is set up for {allowed} replicate(s); this is number "
            f"{row.replicate}. Raise the replicate count on the Setup tab to keep it."
        )


def to_number(raw) -> tuple[float | None, str]:
    """Parse a spreadsheet cell into a number, or explain why it is not one."""
    if raw is None or isinstance(raw, bool):
        return None, "This cell is empty." if raw is None else "This cell holds true/false."
    if isinstance(raw, (int, float)):
        return float(raw), ""
    text = str(raw).strip()
    if not text:
        return None, "This cell is empty."
    if _LESS_OR_MORE.match(text):
        return None, (
            f"\"{text}\" is a limit rather than a reading - it cannot be averaged. "
            "Type the figure you want recorded instead."
        )
    cleaned = text.replace(" ", "").replace("\u00a0", "")
    # A comma is a thousands separator in "2,450" and a decimal point in "2,45".
    # Where both separators appear, whichever comes last is the decimal one, so
    # 1.234,5 and 1,234.5 both end up as 1234.5.
    if "," in cleaned and "." in cleaned:
        cleaned = (cleaned.replace(".", "").replace(",", ".")
                   if cleaned.rfind(",") > cleaned.rfind(".")
                   else cleaned.replace(",", ""))
    elif _THOUSANDS.match(cleaned):
        cleaned = cleaned.replace(",", "")
    else:
        cleaned = cleaned.replace(",", ".")
    try:
        return float(cleaned), ""
    except ValueError:
        return None, f"\"{text}\" is not a number."


# -------------------------------------------------------------------- writing

@dataclass
class Applied:
    """What an import actually did."""

    written: int = 0
    replaced: int = 0
    kept: int = 0    # existing readings left alone because overwrite was off

    def summary(self) -> str:
        parts = [f"{self.written} reading{'s' if self.written != 1 else ''} imported"]
        if self.replaced:
            parts.append(f"{self.replaced} replaced an existing reading")
        if self.kept:
            parts.append(f"{self.kept} left alone because a reading was already there")
        return ", ".join(parts) + "."


def apply(db, run_id: int, plan: ImportPlan, overwrite: bool = False) -> Applied:
    """Write the plan's good rows into a run.

    Readings already in the run are only replaced when ``overwrite`` is set, so
    an import can top up a partly typed round without wiping what is there.
    """
    existing = db.get_run_values(run_id)
    result = Applied()
    for row in plan.ready:
        key = (row.location_id, row.test_id, row.replicate)
        current = existing.get(key)
        if current is not None and current.value is not None:
            if not overwrite:
                result.kept += 1
                continue
            result.replaced += 1
        db.set_value(run_id, row.location_id, row.test_id, row.replicate, row.value, row.note)
        result.written += 1
    return result
