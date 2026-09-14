"""Turn stored measurements into CSV and Excel files.

Three shapes are produced, all built from the same flat replicate rows:

* **raw**     - one row per replicate, the full audit trail.
* **summary** - one row per run/location/test with n, mean, SD, %RSD, range.
* **wide**    - one row per run/location, one column per test (mean values),
                which is the layout most people want for plotting a trend.
"""

from __future__ import annotations

import csv
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence

from .models import LongRow, format_number
from .stats import spec_status, summarize

RAW_HEADERS = [
    "Date", "Time", "Run", "Operator", "Location", "Location code", "Test", "Test code",
    "Unit", "Replicate", "Value", "Lower limit", "Upper limit", "Status", "Note", "Run notes",
]

SUMMARY_HEADERS = [
    "Date", "Time", "Run", "Operator", "Location", "Test", "Unit", "N", "Mean", "SD",
    "%RSD", "Min", "Max", "Range", "Lower limit", "Upper limit", "Status",
]


class ExportError(RuntimeError):
    """Raised when a file cannot be written or a format is unavailable."""


# ------------------------------------------------------------------ grouping

def group_replicates(rows: Iterable[LongRow]) -> "OrderedDict[tuple, list[LongRow]]":
    """Group replicate rows by run, location and test, keeping query order."""
    groups: OrderedDict[tuple, list[LongRow]] = OrderedDict()
    for row in rows:
        groups.setdefault((row.run_id, row.location, row.test), []).append(row)
    return groups


# -------------------------------------------------------------- table shapes

def raw_table(rows: Sequence[LongRow]) -> list[list]:
    table = []
    for row in rows:
        table.append([
            row.run_date, row.run_time, row.run_label, row.operator, row.location,
            row.location_code, row.test, row.test_code, row.unit, row.replicate, row.value,
            row.lower_limit, row.upper_limit,
            _status_text(spec_status(row.value, row.lower_limit, row.upper_limit)),
            row.note, row.run_notes,
        ])
    return table


def summary_table(rows: Sequence[LongRow]) -> list[list]:
    table = []
    for group in group_replicates(rows).values():
        first = group[0]
        stats = summarize([item.value for item in group])
        table.append([
            first.run_date, first.run_time, first.run_label, first.operator, first.location,
            first.test, first.unit, stats.n,
            _round(stats.mean), _round(stats.sd), _round(stats.rsd, 2),
            _round(stats.minimum), _round(stats.maximum), _round(stats.span),
            first.lower_limit, first.upper_limit,
            _status_text(spec_status(stats.mean, first.lower_limit, first.upper_limit)),
        ])
    return table


def wide_table(rows: Sequence[LongRow]) -> tuple[list[str], list[list]]:
    """One row per run/location; one mean value column per test."""
    tests: list[str] = []
    for row in rows:
        if row.test not in tests:
            tests.append(row.test)

    units = {row.test: row.unit for row in rows}
    headers = ["Date", "Time", "Run", "Operator", "Location"] + [
        f"{test} ({units[test]})" if units.get(test) else test for test in tests
    ]

    buckets: OrderedDict[tuple, dict] = OrderedDict()
    for row in rows:
        key = (row.run_id, row.location)
        entry = buckets.setdefault(key, {"meta": row, "values": {}})
        entry["values"].setdefault(row.test, []).append(row.value)

    table = []
    for entry in buckets.values():
        meta: LongRow = entry["meta"]
        line = [meta.run_date, meta.run_time, meta.run_label, meta.operator, meta.location]
        for test in tests:
            line.append(_round(summarize(entry["values"].get(test, [])).mean))
        table.append(line)
    return headers, table


def setup_table(locations, tests) -> tuple[list[str], list[list]]:
    headers = ["Type", "Name", "Code", "Unit", "Decimals", "Replicates",
               "Lower limit", "Upper limit", "Active"]
    table = []
    for location in locations:
        table.append(["Location", location.name, location.code, "", "", "", "", "",
                      "Yes" if location.active else "No"])
    for test in tests:
        table.append(["Test", test.name, test.code, test.unit, test.decimals, test.replicates,
                      test.lower_limit, test.upper_limit, "Yes" if test.active else "No"])
    return headers, table


# --------------------------------------------------------------- CSV writing

def write_csv(path: str | Path, headers: Sequence[str], table: Sequence[Sequence]) -> Path:
    target = Path(path)
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        # utf-8-sig so Excel on Windows opens accented characters correctly.
        with open(target, "w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            writer.writerows(_blank_none(row) for row in table)
    except OSError as error:
        raise ExportError(f"Could not write {target}: {error}") from error
    return target


def export_raw_csv(path, rows: Sequence[LongRow]) -> Path:
    return write_csv(path, RAW_HEADERS, raw_table(rows))


def export_summary_csv(path, rows: Sequence[LongRow]) -> Path:
    return write_csv(path, SUMMARY_HEADERS, summary_table(rows))


def export_wide_csv(path, rows: Sequence[LongRow]) -> Path:
    headers, table = wide_table(rows)
    return write_csv(path, headers, table)


# ------------------------------------------------------------- Excel writing

def excel_available() -> bool:
    """True when this build can write .xlsx files."""
    import importlib.util

    return importlib.util.find_spec("openpyxl") is not None


def export_excel(path, rows: Sequence[LongRow], locations=(), tests=(),
                 title: str = "MeasureLog export") -> Path:
    """Write a multi-sheet workbook: Raw data, Summary, Wide means, Setup."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as error:  # pragma: no cover - depends on the build
        raise ExportError(
            "Excel export needs the openpyxl package. Use CSV export instead, "
            "or install openpyxl and rebuild the app."
        ) from error

    target = Path(path)
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
    except OSError as error:
        raise ExportError(f"Could not create the folder for {target}: {error}") from error

    workbook = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")
    alert_fill = PatternFill("solid", fgColor="FFC7CE")

    wide_headers, wide_rows = wide_table(rows)
    setup_headers, setup_rows = setup_table(locations, tests)
    sheets = [
        ("Raw data", RAW_HEADERS, raw_table(rows), "Status"),
        ("Summary", SUMMARY_HEADERS, summary_table(rows), "Status"),
        ("Wide means", wide_headers, wide_rows, None),
        ("Setup", setup_headers, setup_rows, None),
    ]

    first = True
    for name, headers, table, status_column in sheets:
        sheet = workbook.active if first else workbook.create_sheet()
        sheet.title = name
        first = False

        sheet.append(list(headers))
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        status_index = headers.index(status_column) if status_column in headers else None
        for row in table:
            sheet.append(list(row))
            if status_index is not None and str(row[status_index]).startswith(("Below", "Above")):
                for cell in sheet[sheet.max_row]:
                    cell.fill = alert_fill

        sheet.freeze_panes = "A2"
        for index, header in enumerate(headers, start=1):
            width = max(len(str(header)) + 2, 10)
            for row in table[:200]:
                width = max(width, min(40, len(str(row[index - 1] if row[index - 1] is not None else "")) + 2))
            sheet.column_dimensions[get_column_letter(index)].width = width

    workbook.properties.title = title
    workbook.properties.created = datetime.now()
    try:
        workbook.save(target)
    except OSError as error:
        raise ExportError(
            f"Could not write {target}. If the file is open in Excel, close it and try again."
        ) from error
    return target


# -------------------------------------------------------------------- shared

def default_filename(prefix: str, extension: str, date_from: str = "", date_to: str = "") -> str:
    span = ""
    if date_from and date_to:
        span = f"_{date_from}_to_{date_to}" if date_from != date_to else f"_{date_from}"
    elif date_from:
        span = f"_from_{date_from}"
    return f"{prefix}{span}_{datetime.now():%Y%m%d-%H%M}.{extension.lstrip('.')}"


def _status_text(status: str) -> str:
    return {"ok": "In spec", "low": "Below limit", "high": "Above limit"}.get(status, "")


def _round(value: float | None, digits: int = 4) -> float | None:
    if value is None:
        return None
    return round(float(value), digits)


def _blank_none(row: Sequence) -> list:
    return ["" if value is None else value for value in row]


def describe(rows: Sequence[LongRow]) -> str:
    """Short human summary of what an export will contain."""
    if not rows:
        return "No measurements match the current filter."
    runs = len({row.run_id for row in rows})
    locations = len({row.location for row in rows})
    tests = len({row.test for row in rows})
    dates = sorted({row.run_date for row in rows})
    span = dates[0] if len(dates) == 1 else f"{dates[0]} to {dates[-1]}"
    return (
        f"{len(rows)} measurements | {runs} run(s) | {locations} location(s) | "
        f"{tests} test(s) | {span}"
    )


__all__ = [
    "ExportError", "RAW_HEADERS", "SUMMARY_HEADERS", "default_filename", "describe",
    "excel_available", "export_excel", "export_raw_csv", "export_summary_csv",
    "export_wide_csv", "group_replicates", "raw_table", "setup_table", "summary_table",
    "wide_table", "write_csv", "format_number",
]
